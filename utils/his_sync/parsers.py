import csv
from io import StringIO, BytesIO
from openpyxl import load_workbook

# List of expected 36 HIS Columns
EXPECTED_COLUMNS = [
    "UHID", "PATIENTNAME", "IPNUMBER", "GENDER", "AGE", "ADMITTEDDATE",
    "PLANINGDATE", "CLOSINGDATE", "PRIMARYDOCTOR", "NURSECHECKEDOUTTIME",
    "ADMITING_DOCTOR", "ADMITING_DOCTOR_DEPT", "TREATING_DOCTOR", "TREATING_DOCTOR_DEPT",
    "INTERIMDATE", "DISCHARGE_BILLDATE", "BED_OCCUPIED", "SPECIALIZATION",
    "SPECIALIZATIONDESC", "BEDNO", "BEDTYPE", "BILLABLEBED", "BEDNUMBER",
    "WARDNAME", "COMPANYNAME", "REFRALDOCTOR", "CURRENTSTATUS", "REASONANDREMARKS",
    "CONTACTNO", "ADDRESS1", "COUNTRYNAME", "STATENAME", "CITYNAME",
    "DISTRICTNAME", "CREATEDBY", "LOCATIONID"
]

def normalize_column_name(header):
    """Normalize header name by removing spaces, underscores and converting to uppercase."""
    if not header:
        return ""
    return str(header).strip().upper().replace(" ", "").replace("_", "")

# Mapping normalized expected column names to canonical names
COLUMN_MAP = {normalize_column_name(col): col for col in EXPECTED_COLUMNS}

def _parse_csv(bytes_data):
    try:
        content = bytes_data.decode('utf-8-sig', errors='replace')
    except Exception:
        content = bytes_data.decode('latin-1', errors='replace')
    
    stream = StringIO(content)
    reader = csv.reader(stream)
    raw_headers = next(reader, None)
    if not raw_headers:
        return []

    header_index_map = {}
    for idx, h in enumerate(raw_headers):
        norm_h = normalize_column_name(h)
        if norm_h in COLUMN_MAP:
            header_index_map[idx] = COLUMN_MAP[norm_h]

    csv_rows = []
    for row_idx, row in enumerate(reader, start=2):
        if not row or not any(row):
            continue
        row_dict = {}
        for col_idx, canonical_col in header_index_map.items():
            if col_idx < len(row):
                row_dict[canonical_col] = str(row[col_idx]).strip() if row[col_idx] is not None else ""
        csv_rows.append(row_dict)
    return csv_rows

def _parse_excel(bytes_data):
    wb = load_workbook(filename=BytesIO(bytes_data), data_only=True)
    ws = wb.active
    raw_headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]

    header_index_map = {}
    for idx, h in enumerate(raw_headers):
        if h is not None:
            norm_h = normalize_column_name(h)
            if norm_h in COLUMN_MAP:
                header_index_map[idx] = COLUMN_MAP[norm_h]

    excel_rows = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not any(row):
            continue
        row_dict = {}
        for col_idx, canonical_col in header_index_map.items():
            if col_idx < len(row):
                val = row[col_idx]
                row_dict[canonical_col] = str(val).strip() if val is not None else ""
        excel_rows.append(row_dict)
    return excel_rows

def parse_his_file(file_name, file_bytes):
    """
    Parse Excel (.xlsx) or CSV (.csv) file into a list of row dicts.
    Matches columns by column name tolerating case and whitespace differences.
    Robust against file extension mismatches (e.g. CSV renamed as .xlsx).
    """
    file_name_lower = file_name.lower()

    if file_name_lower.endswith(('.xlsx', '.xls')):
        try:
            return _parse_excel(file_bytes)
        except Exception:
            try:
                return _parse_csv(file_bytes)
            except Exception:
                return []
    else:
        try:
            return _parse_csv(file_bytes)
        except Exception:
            try:
                return _parse_excel(file_bytes)
            except Exception:
                return []
